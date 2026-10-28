from fastapi import FastAPI, Form, HTTPException, Depends
from fastapi.responses import HTMLResponse, FileResponse, RedirectResponse
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from pathlib import Path
from docx import Document
import sqlite3, subprocess, os, secrets, hmac, uuid, json, html
from datetime import datetime

BASE = Path(__file__).resolve().parent
# Persistent storage path on cloud hosts can be set with DATA_DIR (e.g. /var/data).
DATA_DIR = Path(os.getenv('DATA_DIR', str(BASE))).resolve()
TEMPLATE = BASE / 'templates' / 'visa_template.docx'
OUT = DATA_DIR / 'generated'
DB = DATA_DIR / 'visa_applications.sqlite3'
OUT.mkdir(parents=True, exist_ok=True)
app = FastAPI(title='Visa Application Form Generator - Starter')
security = HTTPBasic()

@app.get('/health')
def health():
    return {'status': 'ok'}

FIELDS = [
 ('applicant_name','পূর্ণ নাম','{{APPLICANT_NAME}}'),
 ('father_name','পিতার নাম','{{FATHER_NAME}}'),
 ('mother_name','মায়ের নাম','{{MOTHER_NAME}}'),
 ('date_of_birth','জন্মতারিখ (DD/MM/YYYY)','{{DATE_OF_BIRTH}}'),
 ('place_of_birth','জন্মস্থান','{{PLACE_OF_BIRTH}}'),
 ('marital_status','বৈবাহিক অবস্থা','{{MARITAL_STATUS}}'),
 ('religion','ধর্ম','{{RELIGION}}'),
 ('profession_ar','পেশা (আরবি)','{{PROFESSION_AR}}'),
 ('profession_en','পেশা (ইংরেজি)','{{PROFESSION_EN}}'),
 ('passport_issue_date','পাসপোর্ট ইস্যুর তারিখ (DD/MM/YYYY)','{{PASSPORT_ISSUE_DATE}}'),
 ('passport_expiry_date','পাসপোর্টের মেয়াদ শেষ (DD/MM/YYYY)','{{PASSPORT_EXPIRY_DATE}}'),
 ('passport_number','পাসপোর্ট নম্বর','{{PASSPORT_NUMBER}}'),
]

def auth(credentials: HTTPBasicCredentials = Depends(security)):
    username = os.getenv('VISA_APP_USERNAME')
    password = os.getenv('VISA_APP_PASSWORD')
    if not username or not password:
        raise HTTPException(status_code=503, detail='Set VISA_APP_USERNAME and VISA_APP_PASSWORD before running.')
    ok = hmac.compare_digest(credentials.username, username) and hmac.compare_digest(credentials.password, password)
    if not ok:
        raise HTTPException(status_code=401, detail='Authentication required', headers={'WWW-Authenticate':'Basic'})
    return credentials.username

def init_db():
    with sqlite3.connect(DB) as con:
        con.execute('''CREATE TABLE IF NOT EXISTS passengers (
            id INTEGER PRIMARY KEY AUTOINCREMENT, applicant_name TEXT NOT NULL, passport_number TEXT NOT NULL,
            created_at TEXT NOT NULL, UNIQUE(passport_number))''')
        con.execute('''CREATE TABLE IF NOT EXISTS applications (
            id TEXT PRIMARY KEY, passenger_id INTEGER NOT NULL, payload_json TEXT NOT NULL,
            docx_path TEXT NOT NULL, pdf_path TEXT NOT NULL, created_at TEXT NOT NULL,
            FOREIGN KEY(passenger_id) REFERENCES passengers(id))''')
init_db()

def replace_in_paragraph(paragraph, replacements):
    # Placeholders are intentionally kept inside one run in the configured template.
    for run in paragraph.runs:
        for token, value in replacements.items():
            if token in run.text:
                run.text = run.text.replace(token, value)

def replace_in_doc(doc, replacements):
    for p in doc.paragraphs: replace_in_paragraph(p, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs: replace_in_paragraph(p, replacements)

def page(title, body):
    return HTMLResponse(f'''<!doctype html><html lang="bn"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title>
    <style>body{{font-family:Arial,'Noto Sans Bengali',sans-serif;max-width:850px;margin:30px auto;padding:0 16px;color:#202124}}h1{{font-size:26px}}form{{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:14px}}label{{display:block;font-weight:600}}input{{box-sizing:border-box;width:100%;padding:10px;margin-top:6px;border:1px solid #bbb;border-radius:6px}}button,a.btn{{background:#174ea6;color:white;border:0;padding:11px 16px;border-radius:6px;cursor:pointer;text-decoration:none;display:inline-block}}.note{{background:#fff7d6;padding:12px;border-radius:6px;line-height:1.5}}.wide{{grid-column:1/-1}}</style><body>{body}</body></html>''')

@app.get('/', response_class=HTMLResponse)
def home(user=Depends(auth)):
    inputs=''.join(f'<label>{html.escape(label)}<input name="{key}" required="required" autocomplete="off"></label>' for key,label,_ in FIELDS)
    body=f'''<h1>ভিসা আবেদন ফর্ম জেনারেটর</h1><p class="note">এটি একটি প্রাথমিক সংস্করণ। “For Official use only” অংশে ভিসা নম্বর বা অনুমোদন নম্বর তৈরি/পরিবর্তন করা হয় না; ওই অংশ অনুমোদিত কর্তৃপক্ষের জন্য সংরক্ষিত।</p>
    <h2>নতুন আবেদন</h2><form method="post" action="/generate">{inputs}<div class="wide"><button type="submit">তথ্য সংরক্ষণ ও PDF তৈরি</button> <a class="btn" href="/history">আবেদন ইতিহাস</a></div></form>'''
    return page('নতুন আবেদন',body)

@app.post('/generate')
def generate(user=Depends(auth), applicant_name:str=Form(...), father_name:str=Form(...), mother_name:str=Form(...), date_of_birth:str=Form(...), place_of_birth:str=Form(...), marital_status:str=Form(...), religion:str=Form(...), profession_ar:str=Form(...), profession_en:str=Form(...), passport_issue_date:str=Form(...), passport_expiry_date:str=Form(...), passport_number:str=Form(...)):
    data=locals()
    payload={k:data[k].strip() for k,_,_ in FIELDS}
    if any(not v for v in payload.values()): raise HTTPException(400,'সব ফিল্ড পূরণ করুন।')
    app_id=uuid.uuid4().hex
    work=OUT/app_id; work.mkdir()
    docx_path=work/'application.docx'; pdf_path=work/'application.pdf'
    doc=Document(TEMPLATE)
    replacements={token:payload[key] for key,_,token in FIELDS}
    replace_in_doc(doc,replacements)
    doc.save(docx_path)
    try:
        subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo_'+app_id,'--headless','--convert-to','pdf','--outdir',str(work),str(docx_path)],check=True,capture_output=True,timeout=60)
    except Exception as e:
        raise HTTPException(500,'PDF তৈরি করা যায়নি। সার্ভারে LibreOffice ইনস্টল আছে কি না পরীক্ষা করুন।') from e
    now=datetime.utcnow().isoformat(timespec='seconds')+'Z'
    with sqlite3.connect(DB) as con:
        con.execute('PRAGMA foreign_keys=ON')
        row=con.execute('SELECT id FROM passengers WHERE passport_number=?',(payload['passport_number'],)).fetchone()
        if row:
            passenger_id=row[0]
            con.execute('UPDATE passengers SET applicant_name=?,created_at=? WHERE id=?',(payload['applicant_name'],now,passenger_id))
        else:
            cur=con.execute('INSERT INTO passengers(applicant_name,passport_number,created_at) VALUES(?,?,?)',(payload['applicant_name'],payload['passport_number'],now)); passenger_id=cur.lastrowid
        con.execute('INSERT INTO applications(id,passenger_id,payload_json,docx_path,pdf_path,created_at) VALUES(?,?,?,?,?,?)',(app_id,passenger_id,json.dumps(payload,ensure_ascii=False),str(docx_path),str(pdf_path),now))
    return RedirectResponse(f'/result/{app_id}',status_code=303)

@app.get('/result/{app_id}',response_class=HTMLResponse)
def result(app_id:str,user=Depends(auth)):
    with sqlite3.connect(DB) as con:
        row=con.execute('SELECT pdf_path FROM applications WHERE id=?',(app_id,)).fetchone()
    if not row: raise HTTPException(404,'আবেদন পাওয়া যায়নি')
    body=f'<h1>PDF তৈরি হয়েছে</h1><p>আবেদনটি ডেটাবেজে সংরক্ষিত হয়েছে।</p><p><a class="btn" href="/download/{app_id}">PDF ডাউনলোড</a> <a class="btn" href="/history">আবেদন ইতিহাস</a></p><iframe src="/download/{app_id}" style="width:100%;height:75vh;border:1px solid #ddd"></iframe>'
    return page('PDF তৈরি',body)

@app.get('/download/{app_id}')
def download(app_id:str,user=Depends(auth)):
    with sqlite3.connect(DB) as con: row=con.execute('SELECT pdf_path FROM applications WHERE id=?',(app_id,)).fetchone()
    if not row or not Path(row[0]).exists(): raise HTTPException(404,'PDF পাওয়া যায়নি')
    return FileResponse(row[0],media_type='application/pdf',filename=f'visa-application-{app_id[:8]}.pdf')

@app.get('/history',response_class=HTMLResponse)
def history(user=Depends(auth)):
    with sqlite3.connect(DB) as con:
        rows=con.execute('SELECT a.id,p.applicant_name,p.passport_number,a.created_at FROM applications a JOIN passengers p ON p.id=a.passenger_id ORDER BY a.created_at DESC LIMIT 100').fetchall()
    trs=''.join(f'<tr><td>{html.escape(r[1])}</td><td>{html.escape(r[2])}</td><td>{html.escape(r[3])}</td><td><a href="/download/{r[0]}">PDF</a></td></tr>' for r in rows)
    body=f'<h1>আবেদন ইতিহাস</h1><p><a class="btn" href="/">নতুন আবেদন</a></p><table border="1" cellpadding="8" cellspacing="0"><tr><th>নাম</th><th>পাসপোর্ট নম্বর</th><th>তৈরির সময়</th><th>PDF</th></tr>{trs}</table>'
    return page('আবেদন ইতিহাস',body)
